import openpyxl
from pprint import pprint

wb = openpyxl.load_workbook('censuspopdata.xlsx')
sheet = wb['Population by Census Tract']

data = {}

for row in range(2, sheet.max_row + 1):
    s = sheet.cell(row=row, column=2).value 
    c = sheet.cell(row=row, column=3).value
    p = sheet.cell(row=row, column=4).value
    data[row] = {'State': s, 'County': c, 'Population': p}

with open('census2010.py', 'w') as f:
    f.write('data = ')
    pprint(data, stream=f)

